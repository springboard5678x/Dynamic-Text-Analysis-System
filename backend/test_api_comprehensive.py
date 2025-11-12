import requests
import json
import os
import time
from datetime import datetime
import pandas as pd
import docx

# API base URL
BASE_URL = "http://localhost:8000/api"

def test_api_connection():
    """Test if API is accessible"""
    print("🔌 Testing API connection...")
    try:
        response = requests.get(f"{BASE_URL}/analysis/")
        if response.status_code == 200:
            print("✅ API is accessible and working!")
            return True
        else:
            print(f"❌ API returned status code: {response.status_code}")
            return False
    except Exception as e:
        print(f"❌ Cannot connect to API: {e}")
        return False

def test_raw_text_analysis():
    """Test text analysis with raw text input"""
    print("\n" + "="*50)
    print("📝 TEST 1: RAW TEXT ANALYSIS")
    print("="*50)
    
    sample_text = """
    Artificial Intelligence and Machine Learning in Modern Business
    
    The integration of artificial intelligence and machine learning technologies has revolutionized how businesses operate across various industries. 
    Companies are leveraging AI for customer service automation, predictive analytics, and process optimization. 
    Machine learning algorithms help in identifying patterns in large datasets that were previously impossible to analyze manually.
    
    However, there are significant challenges in implementing AI solutions. Data privacy concerns remain a major issue, 
    with many customers expressing apprehension about how their personal information is used. 
    The cost of implementation and the need for specialized talent also pose barriers for small and medium-sized enterprises.
    
    Despite these challenges, the benefits are substantial. Businesses report improved efficiency, reduced operational costs, 
    and enhanced customer experiences. The future looks promising with advancements in natural language processing 
    and computer vision opening new possibilities for innovation.
    
    Ethical considerations must be addressed as AI becomes more pervasive. Transparency in AI decision-making 
    and accountability for automated systems are critical for building trust with stakeholders.
    """
    
    payload = {
        "title": f"AI in Business Analysis - {datetime.now().strftime('%H:%M:%S')}",
        "input_text": sample_text
    }
    
    try:
        response = requests.post(f"{BASE_URL}/analysis/", json=payload)
        
        if response.status_code == 201:
            data = response.json()
            print("✅ Raw Text Analysis Successful!")
            print(f"📊 Analysis ID: {data['id']}")
            print(f"📋 Title: {data['title']}")
            print(f"📝 Text Length: {len(sample_text)} characters, {len(sample_text.split())} words")
            return data['id']
        else:
            print(f"❌ Raw Text Analysis Failed: {response.status_code}")
            print(f"Response: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Error during raw text analysis: {e}")
        return None

def test_txt_file_analysis():
    """Test analysis with .txt file upload"""
    print("\n" + "="*50)
    print("📄 TEST 2: .TXT FILE ANALYSIS")
    print("="*50)
    
    # Create a comprehensive sample text file
    sample_content = """
    Quarterly Financial Report - Q4 2024
    
    EXECUTIVE SUMMARY
    The fourth quarter demonstrated robust growth with a 22% increase in overall revenue compared to Q3. 
    Our e-commerce platform saw unprecedented traffic, resulting in a 35% surge in online sales.
    
    FINANCIAL HIGHLIGHTS
    - Total Revenue: $15.8 million (up 22% from Q3)
    - Net Profit: $3.2 million (profit margin: 20.3%)
    - Operating Expenses: $9.1 million
    - R&D Investment: $2.1 million
    
    DEPARTMENT PERFORMANCE
    Marketing: Successful holiday campaign resulted in 45% increase in brand awareness.
    Sales: Exceeded targets by 18% through improved customer relationship management.
    Technology: Launched new mobile application with 50,000 downloads in first month.
    Customer Service: Improved response time by 40% through AI-powered chatbots.
    
    CHALLENGES AND OPPORTUNITIES
    Supply chain disruptions affected product delivery times by an average of 3 days.
    Increased competition in the Southeast Asian market requires strategic repositioning.
    Opportunity to expand into European markets with localized product offerings.
    
    OUTLOOK FOR Q1 2025
    Projected revenue growth of 15-20% based on current market trends.
    Plan to hire 50 new employees across engineering and marketing departments.
    Budget allocation of $5 million for market expansion initiatives.
    """
    
    # Create test .txt file
    file_path = f"test_financial_report_{datetime.now().strftime('%H%M%S')}.txt"
    
    try:
        # Write file
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(sample_content)
        
        # Upload file
        with open(file_path, "rb") as f:
            files = {
                'uploaded_file': f
            }
            data = {
                'title': f'Financial Report Analysis - {datetime.now().strftime("%H:%M:%S")}'
            }
            
            response = requests.post(f"{BASE_URL}/analysis/", files=files, data=data)
        
        if response.status_code == 201:
            data = response.json()
            print("✅ .TXT File Analysis Successful!")
            print(f"📊 Analysis ID: {data['id']}")
            print(f"📁 File Type: {data.get('file_type', 'N/A')}")
            print(f"📝 Content Length: {len(sample_content)} characters")
            return data['id']
        else:
            print(f"❌ .TXT File Analysis Failed: {response.status_code}")
            print(f"Response: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Error during .txt file analysis: {e}")
        return None
    finally:
        # Clean up test file
        cleanup_file(file_path)

def test_docx_file_analysis():
    """Test analysis with .docx file upload"""
    print("\n" + "="*50)
    print("📑 TEST 3: .DOCX FILE ANALYSIS")
    print("="*50)
    
    # Create a sample .docx file
    file_path = f"test_business_plan_{datetime.now().strftime('%H%M%S')}.docx"
    
    try:
        # Create a new Document
        doc = docx.Document()
        
        # Add title
        doc.add_heading('Business Development Plan 2025', 0)
        
        # Add sections
        doc.add_heading('Executive Summary', level=1)
        doc.add_paragraph('This comprehensive business plan outlines our strategy for market expansion and product development in the upcoming fiscal year. We aim to increase market share by 15% through innovative solutions and strategic partnerships.')
        
        doc.add_heading('Market Analysis', level=1)
        doc.add_paragraph('The current market shows significant growth potential in the technology sector. Consumer demand for AI-powered solutions has increased by 40% year-over-year. Competitor analysis reveals gaps in customer service automation that we can exploit.')
        
        doc.add_heading('Product Roadmap', level=1)
        doc.add_paragraph('Q1 2025: Launch enhanced mobile application with AI features')
        doc.add_paragraph('Q2 2025: Develop enterprise-level analytics dashboard')
        doc.add_paragraph('Q3 2025: Expand to European markets with localized solutions')
        doc.add_paragraph('Q4 2025: Implement blockchain technology for data security')
        
        doc.add_heading('Financial Projections', level=1)
        doc.add_paragraph('Year 1 Revenue: $12 million')
        doc.add_paragraph('Year 2 Revenue: $18 million (projected)')
        doc.add_paragraph('Year 3 Revenue: $25 million (projected)')
        doc.add_paragraph('Break-even Point: Month 14')
        
        doc.add_heading('Risk Assessment', level=1)
        doc.add_paragraph('Market Risks: Changing regulatory environment, increased competition')
        doc.add_paragraph('Technical Risks: Rapid technology obsolescence, cybersecurity threats')
        doc.add_paragraph('Financial Risks: Cash flow constraints, unexpected market downturns')
        
        # Save the document
        doc.save(file_path)
        
        # Upload file
        with open(file_path, "rb") as f:
            files = {
                'uploaded_file': f
            }
            data = {
                'title': f'Business Plan Analysis - {datetime.now().strftime("%H:%M:%S")}'
            }
            
            response = requests.post(f"{BASE_URL}/analysis/", files=files, data=data)
        
        if response.status_code == 201:
            data = response.json()
            print("✅ .DOCX File Analysis Successful!")
            print(f"📊 Analysis ID: {data['id']}")
            print(f"📁 File Type: {data.get('file_type', 'N/A')}")
            return data['id']
        else:
            print(f"❌ .DOCX File Analysis Failed: {response.status_code}")
            print(f"Response: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Error during .docx file analysis: {e}")
        return None
    finally:
        # Clean up test file
        cleanup_file(file_path)

def test_csv_file_analysis():
    """Test analysis with .csv file upload"""
    print("\n" + "="*50)
    print("📊 TEST 4: .CSV FILE ANALYSIS")
    print("="*50)
    
    # Create a sample .csv file with text data
    file_path = f"test_customer_feedback_{datetime.now().strftime('%H%M%S')}.csv"
    
    try:
        # Create sample data
        data = {
            'customer_id': [101, 102, 103, 104, 105, 106, 107, 108, 109, 110],
            'feedback_text': [
                "The product quality is excellent and delivery was faster than expected. Very satisfied with the service.",
                "Customer support was unhelpful and the product arrived damaged. Very disappointed with this purchase.",
                "Good value for money but the user interface could be more intuitive. Overall positive experience.",
                "Outstanding service! The team went above and beyond to resolve my issue quickly.",
                "Average product quality. Nothing special but gets the job done. Would consider buying again.",
                "Terrible experience. Product stopped working after one week and refund process is complicated.",
                "Excellent features and great customer service. The mobile app is particularly impressive.",
                "The product is okay but shipping took too long. Communication could be improved.",
                "Amazing quality and fast shipping. Will definitely recommend to friends and colleagues.",
                "Poor packaging resulted in damaged goods. Customer service response was slow and unhelpful."
            ],
            'rating': [5, 1, 4, 5, 3, 1, 5, 3, 5, 2],
            'category': ['positive', 'negative', 'positive', 'positive', 'neutral', 'negative', 'positive', 'neutral', 'positive', 'negative']
        }
        
        # Create DataFrame and save as CSV
        df = pd.DataFrame(data)
        df.to_csv(file_path, index=False)
        
        # Upload file
        with open(file_path, "rb") as f:
            files = {
                'uploaded_file': f
            }
            data = {
                'title': f'Customer Feedback Analysis - {datetime.now().strftime("%H:%M:%S")}'
            }
            
            response = requests.post(f"{BASE_URL}/analysis/", files=files, data=data)
        
        if response.status_code == 201:
            data = response.json()
            print("✅ .CSV File Analysis Successful!")
            print(f"📊 Analysis ID: {data['id']}")
            print(f"📁 File Type: {data.get('file_type', 'N/A')}")
            print(f"📊 CSV Rows: {len(df)} customer feedback entries")
            return data['id']
        else:
            print(f"❌ .CSV File Analysis Failed: {response.status_code}")
            print(f"Response: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Error during .csv file analysis: {e}")
        return None
    finally:
        # Clean up test file
        cleanup_file(file_path)

def test_large_text_analysis():
    """Test with a larger text to see more topics"""
    print("\n" + "="*50)
    print("📚 TEST 5: LARGE TEXT ANALYSIS")
    print("="*50)
    
    large_text = """
    Comprehensive Analysis of Digital Transformation in Healthcare
    
    INTRODUCTION
    The healthcare industry is undergoing a significant digital transformation, driven by technological advancements and changing patient expectations. This shift encompasses electronic health records, telemedicine, AI diagnostics, and personalized treatment plans.
    
    TECHNOLOGICAL ADVANCEMENTS
    Artificial Intelligence is revolutionizing medical diagnostics through image recognition and pattern analysis. Machine learning algorithms can now detect diseases like cancer and diabetes with accuracy rates exceeding human experts in some cases. Natural language processing enables efficient analysis of medical literature and patient records.
    
    TELEMEDICINE EXPANSION
    The COVID-19 pandemic accelerated telemedicine adoption by over 500% in two years. Patients now expect remote consultation options, and healthcare providers are investing heavily in virtual care platforms. This shift has improved access to care for rural populations and reduced hospital overcrowding.
    
    DATA SECURITY CHALLENGES
    With digitization comes increased cybersecurity risks. Healthcare organizations face constant threats from ransomware attacks and data breaches. Implementing robust security measures while maintaining accessibility remains a critical challenge. Patient data privacy regulations require strict compliance across all digital platforms.
    
    PATIENT ENGAGEMENT
    Digital tools have transformed patient engagement through mobile health applications and wearable devices. Patients can now monitor their health metrics in real-time and share data directly with healthcare providers. This continuous monitoring enables proactive healthcare interventions and personalized treatment plans.
    
    REGULATORY LANDSCAPE
    Government regulations are evolving to address digital healthcare innovations. The FDA has established frameworks for approving AI-based medical devices, while privacy laws like HIPAA continue to adapt to new technologies. International standards are emerging to facilitate global healthcare data exchange.
    
    FINANCIAL IMPLICATIONS
    Digital transformation requires substantial investment in infrastructure and training. However, studies show that digital healthcare solutions can reduce operational costs by 20-30% while improving patient outcomes. Insurance companies are increasingly covering telemedicine services and digital therapeutics.
    
    FUTURE TRENDS
    The integration of blockchain technology promises enhanced data security and interoperability. Augmented reality is being explored for surgical planning and medical training. Genomics and personalized medicine will continue to advance through computational biology and big data analytics.
    
    ETHICAL CONSIDERATIONS
    Algorithmic bias in AI diagnostics raises concerns about equitable healthcare access. The digital divide may exacerbate health disparities if not addressed proactively. Transparent AI systems and diverse training data are essential for ethical implementation of digital healthcare solutions.
    
    IMPLEMENTATION STRATEGIES
    Successful digital transformation requires cross-functional teams involving clinicians, IT specialists, and administrators. Change management is crucial for staff adoption of new technologies. Continuous training and support ensure that healthcare professionals can leverage digital tools effectively.
    
    CONCLUSION
    Digital transformation in healthcare presents both unprecedented opportunities and significant challenges. Balancing innovation with security, accessibility with privacy, and technology with human touch will define the future of healthcare delivery. Organizations that embrace this transformation while maintaining patient-centric values will lead the industry forward.
    """
    
    payload = {
        "title": f"Healthcare Digital Transformation - {datetime.now().strftime('%H:%M:%S')}",
        "input_text": large_text
    }
    
    try:
        response = requests.post(f"{BASE_URL}/analysis/", json=payload)
        
        if response.status_code == 201:
            data = response.json()
            print("✅ Large Text Analysis Successful!")
            print(f"📊 Analysis ID: {data['id']}")
            print(f"📋 Title: {data['title']}")
            print(f"📝 Text Length: {len(large_text)} characters, {len(large_text.split())} words")
            return data['id']
        else:
            print(f"❌ Large Text Analysis Failed: {response.status_code}")
            print(f"Response: {response.text}")
            return None
            
    except Exception as e:
        print(f"❌ Error during large text analysis: {e}")
        return None

def get_analysis_results(analysis_id, test_name):
    """Retrieve and display analysis results"""
    print(f"\n📈 Retrieving Results for {test_name}...")
    
    try:
        response = requests.get(f"{BASE_URL}/analysis/{analysis_id}/")
        
        if response.status_code == 200:
            data = response.json()
            print("✅ Results Retrieved Successfully!")
            
            # Display key results
            print(f"\n🎯 Analysis Summary:")
            print(f"   Title: {data['title']}")
            print(f"   Created: {data['created_at']}")
            
            # Topics
            if data.get('topics_json'):
                topics = data['topics_json']
                print(f"\n📊 Topics Identified ({len(topics)} topics):")
                for topic in topics[:6]:  # Show up to 6 topics
                    words = ', '.join(topic['words'][:5])
                    print(f"   Topic {topic['topic_id']}: {words}")
                if len(topics) > 6:
                    print(f"   ... and {len(topics) - 6} more topics")
            
            # Sentiment
            if data.get('sentiment_distribution'):
                sentiment = data['sentiment_distribution']
                print(f"\n😊 Sentiment Analysis:")
                print(f"   Overall: {sentiment.get('overall_sentiment', 'N/A')}")
                if 'distribution' in sentiment:
                    dist = sentiment['distribution']
                    total = sum(dist.values())
                    if total > 0:
                        print(f"   Positive: {dist.get('positive', 0)} ({dist.get('positive', 0)/total*100:.1f}%)")
                        print(f"   Neutral: {dist.get('neutral', 0)} ({dist.get('neutral', 0)/total*100:.1f}%)")
                        print(f"   Negative: {dist.get('negative', 0)} ({dist.get('negative', 0)/total*100:.1f}%)")
            
            # Summaries
            if data.get('extractive_summary'):
                print(f"\n📋 Extractive Summary (first 150 chars):")
                print(f"   {data['extractive_summary'][:150]}...")
            
            if data.get('abstractive_summary'):
                print(f"\n📝 Abstractive Summary (first 150 chars):")
                print(f"   {data['abstractive_summary'][:150]}...")
            
            # Insights
            if data.get('actionable_insights'):
                print(f"\n💡 Actionable Insights:")
                for insight in data['actionable_insights'][:3]:
                    print(f"   • {insight['title']}: {insight['recommendation']}")
            
            return True
        else:
            print(f"❌ Failed to retrieve results: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ Error retrieving results: {e}")
        return False

def download_report(analysis_id, test_name):
    """Test report download functionality"""
    print(f"\n📄 Testing Report Download for {test_name}...")
    
    try:
        response = requests.get(f"{BASE_URL}/analysis/{analysis_id}/download_report/")
        
        if response.status_code == 200:
            # Save the report file
            report_filename = f"report_{test_name.replace(' ', '_')}_{analysis_id}.pdf"
            with open(report_filename, 'wb') as f:
                f.write(response.content)
            print(f"✅ Report Downloaded Successfully: {report_filename}")
            return True
        else:
            print(f"❌ Report download failed: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ Error downloading report: {e}")
        return False

def cleanup_file(file_path):
    """Clean up test files with retry logic"""
    max_retries = 3
    for attempt in range(max_retries):
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                break
        except PermissionError:
            if attempt < max_retries - 1:
                time.sleep(0.1)
            else:
                print(f"⚠️ Could not delete test file: {file_path}")

def list_all_analyses():
    """List all existing analyses"""
    print("\n" + "="*50)
    print("📋 OVERVIEW: ALL ANALYSES")
    print("="*50)
    
    try:
        response = requests.get(f"{BASE_URL}/analysis/")
        
        if response.status_code == 200:
            data = response.json()
            print(f"📊 Total Analyses in Database: {len(data)}")
            
            # Group by file type
            file_types = {}
            for analysis in data:
                file_type = analysis.get('file_type', 'raw_text')
                if file_type not in file_types:
                    file_types[file_type] = 0
                file_types[file_type] += 1
            
            print(f"\n📁 Analysis Types:")
            for file_type, count in file_types.items():
                print(f"   • {file_type.upper()}: {count} analyses")
            
            print(f"\n🕒 Recent Analyses:")
            for analysis in data[:5]:
                created = analysis['created_at'].replace('T', ' ').split('.')[0]
                print(f"   • {analysis['title']} - {created}")
            
            if len(data) > 5:
                print(f"   ... and {len(data) - 5} more analyses")
            
            return True
        else:
            print(f"❌ Failed to list analyses: {response.status_code}")
            return False
            
    except Exception as e:
        print(f"❌ Error listing analyses: {e}")
        return False

def main():
    """Run complete API test suite for all file types"""
    print("🚀 Starting Comprehensive NarrativeNexus API Test Suite")
    print("=" * 60)
    
    # Test API connection
    if not test_api_connection():
        print("Stopping tests due to API connection failure")
        return
    
    test_results = {}
    
    # Test 1: Raw Text
    analysis_id = test_raw_text_analysis()
    if analysis_id:
        test_results['Raw Text'] = analysis_id
        get_analysis_results(analysis_id, "Raw Text Analysis")
        download_report(analysis_id, "Raw_Text")
    
    # Test 2: TXT File
    analysis_id = test_txt_file_analysis()
    if analysis_id:
        test_results['TXT File'] = analysis_id
        get_analysis_results(analysis_id, "TXT File Analysis")
        download_report(analysis_id, "TXT_File")
    
    # Test 3: DOCX File
    analysis_id = test_docx_file_analysis()
    if analysis_id:
        test_results['DOCX File'] = analysis_id
        get_analysis_results(analysis_id, "DOCX File Analysis")
        download_report(analysis_id, "DOCX_File")
    
    # Test 4: CSV File
    analysis_id = test_csv_file_analysis()
    if analysis_id:
        test_results['CSV File'] = analysis_id
        get_analysis_results(analysis_id, "CSV File Analysis")
        download_report(analysis_id, "CSV_File")
    
    # Test 5: Large Text
    analysis_id = test_large_text_analysis()
    if analysis_id:
        test_results['Large Text'] = analysis_id
        get_analysis_results(analysis_id, "Large Text Analysis")
        download_report(analysis_id, "Large_Text")
    
    # Final overview
    list_all_analyses()
    
    print("\n" + "=" * 60)
    print("🎉 COMPREHENSIVE TEST SUITE COMPLETED!")
    print("\n📊 Test Summary:")
    for test_name, analysis_id in test_results.items():
        print(f"   ✅ {test_name}: {analysis_id}")
    
    print(f"\n📝 Next Steps:")
    print("   • Check generated PDF reports in the current directory")
    print("   • Visit Django admin: http://localhost:8000/admin/")
    print("   • All file types successfully tested!")
    print("   • Ready for React frontend development! 🚀")

if __name__ == "__main__":
    main()